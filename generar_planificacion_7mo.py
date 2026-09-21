# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import copy

def create_7mo_plan():
    doc = docx.Document('LENGUA TERCER  TRIMESTRE CONTINUACION.docx')
    table = doc.tables[0]
    
    # 1. Modificar encabezado de institución y periodo (Row 0)
    for p in table.rows[0].cells[0].paragraphs:
        if '2025' in p.text or 'PLANIFICACI' in p.text:
            p.text = p.text.replace('2025 – 2026', '2026 – 2027').replace('2025 - 2026', '2026 – 2027')
            p.text = p.text.replace('PLANIFICACIÓN MICROCURRICULAR TRIMESTRAL', 'PLANIFICACIÓN MICROCURRICULAR DEL PRIMER TRIMESTRE (8 SEMANAS)')

    # 2. Modificar Datos Informativos (Row 2 y Row 3)
    # Row 2: Docentes
    r2_c1 = table.rows[2].cells[1]
    r2_c1.text = "LIC. KLEVER YANEZ / MSC. MERY DÍAZ / MSC. LUIS ROBLES / MSC. MARÍA ISABEL RIVAS / MSC. LAURA TORRES"
    
    # Row 3: Grado, trimestre y fechas oficiales
    r3 = table.rows[3]
    r3.cells[1].text = "SÉPTIMO"
    r3.cells[7].text = "A, B, C, D Y E (Asignado: 7mo \"C\")"
    r3.cells[12].text = "1er Trimestre"
    r3.cells[17].text = "05/10/2026"
    r3.cells[21].text = "30/11/2026 (Microcurricular: 05/10 al 13/11 · Proyecto: 16/11 al 23/11 · Exámenes: 24/11 al 30/11)"

    # 3. Aprendizaje Disciplinar (Row 5: Objetivos de Séptimo)
    r5_cell = table.rows[5].cells[3]
    r5_cell.text = (
        "O.LL.3.1. Interactuar con diversas expresiones culturales para acceder, participar y apropiarse de la cultura escrita.\n"
        "O.LL.3.3. Comprender discursos orales en diversos contextos de la actividad social y cultural y analizarlos con sentido crítico.\n"
        "O.LL.3.6. Leer de manera autónoma textos no literarios, con fines de recreación, información y aprendizaje, y aplicar estrategias cognitivas de comprensión.\n"
        "O.LL.3.8. Escribir relatos y textos expositivos, descriptivos e instructivos, adecuados a una situación comunicativa determinada para aprender y comunicarse.\n"
        "O.LL.3.11. Seleccionar y disfrutar textos literarios para realizar interpretaciones personales y construir significados compartidos con otros lectores."
    )

    # 4. Semanas de planificación: 6 microcurriculares + 1 proyecto integrador + 1 exámenes = 8 semanas
    # Insertar 2 filas clonando row 14 para semanas 12 y 13 antes de las firmas (row 15)
    r14 = table.rows[14]
    r15 = table.rows[15]
    
    # Clon 1 para semana 12
    clone1 = copy.deepcopy(r14._tr)
    r15._tr.addprevious(clone1)
    
    # Clon 2 para semana 13
    clone2 = copy.deepcopy(r14._tr)
    r15._tr.addprevious(clone2)

    weeks_data = [
        {
            # SEMANA 6 (Semana 1 de la Microcurricular) - ORÍGENES DE LA ESCRITURA (Pág. 12)
            "row_idx": 9,
            "dcd": "LL.3.1.1. Participar en contextos y situaciones que evidencien la funcionalidad de la lengua escrita como herramienta cultural.",
            "tema": "Tema: Lengua y Cultura: Orígenes de la escritura (de las pinturas rupestres a la era digital)\n\nMateriales:\nTexto de Lengua 7mo (Páginas 12 a 17)\nPizarra y marcadores de colores\nCuadernos de trabajo\nGuía Docente Maya/MinEduc pág. 12",
            "estrategias": [
                ("SEMANA 6 (Semana 1 de la Planificación Microcurricular · 05/10 al 09/10/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente solicita a los estudiantes abrir el libro en la página 12 y observar las imágenes iniciales: pinturas rupestres con cazadores y fauna, alfabetos primitivos, jeroglíficos, escriba egipcio, caligrafía del abecedario y un estudiante frente a una computadora moderna.\nPreguntas detonantes de la sección \"Mi experiencia\": ¿Qué conoces sobre las primeras formas de escritura?, ¿cómo ha evolucionado la escritura en el tiempo?, ¿en tu cotidianidad, para qué usas la palabra escrita?, ¿crees que actualmente la palabra escrita tiene más valor que la palabra oral? ¿Por qué?", False, 10),
                ("Inserción Educación Cívica, Ética e Integridad:", True, 10),
                ("Reflexión sobre el valor de la palabra escrita para preservar la memoria, los acuerdos comunitarios, las leyes y los derechos humanos a lo largo de la historia de los pueblos.", False, 10),
                ("DUA", True, 10),
                ("Representación: ilustración paratextual histórica de la pág. 12, lectura en voz alta comentada.\nAcción/expresión: participación oral y lluvia de ideas sobre situaciones donde usamos la escritura hoy.\nMotivación: curiosidad sobre cómo se comunicaban los seres humanos hace miles de años.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("El docente guía la lectura y diálogo de la situación comunicativa inicial: ¿Por qué sintió el ser humano la necesidad de crear un sistema gráfico si ya existía la comunicación oral? Los estudiantes deducen que las palabras orales se las lleva el viento, mientras que la escritura permite registrar cosechas, leyes, transacciones y conocimientos sin que se deformen con el paso de los siglos.", False, 10),
                ("DUA", True, 10),
                ("Representación: preguntas de análisis proyectadas y anotadas en pizarra.\nAcción/expresión: debate moderado en parejas y exposición de conclusiones.\nMotivación: descubrimiento reflexivo de la utilidad social de la escritura.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Lectura analítica guiada de las páginas 13 a 17 del texto escolar sobre la evolución histórica de los soportes y sistemas de escritura:\n1. Arte rupestre y pictogramas.\n2. Escritura cuneiforme en tablillas de arcilla (Sumeria).\n3. Jeroglíficos y el papiro egipcio.\n4. Invención del alfabeto fenicio, griego y latino.\n5. El pergamino, el papel de trapo, la imprenta de Gutenberg y los medios digitales.\nConstrucción en la pizarra de un organizador gráfico (Línea de Tiempo Ilustrada) que los estudiantes completan en sus cuadernos.", False, 10),
                ("DUA", True, 10),
                ("Representación: organizador gráfico cronológico en pizarra con marcadores de colores.\nAcción/expresión: registro ordenado en el cuaderno con glosario de términos (pictograma, ideograma, cuneiforme).\nMotivación: visión histórica integral de la evolución cultural humana.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Los estudiantes resuelven de forma autónoma el taller de aplicación de la página 17 del texto escolar:\n- Descifran el mensaje secreto utilizando el código de sustitución alfabética.\n- Escriben en su cuaderno un párrafo reflexivo respondiendo: \"¿Por qué la lengua escrita es una herramienta cultural imprescindible en mi vida escolar y familiar?\"\nCoevaluación en parejas y socialización. (Nota: Viernes 09 de Octubre: Feriado por la Independencia de Guayaquil).", False, 10),
                ("DUA", True, 10),
                ("Representación: actividad lúdica de criptografía y descifrado de códigos en pág. 17.\nAcción/expresión: redacción breve en el cuaderno y retroalimentación entre pares.\nMotivación: reto lúdico de resolución y afianzamiento del aprendizaje.", False, 10),
            ],
            "indicador": "Reconoce la funcionalidad de la lengua escrita como manifestación cultural y de identidad en diferentes contextos y situaciones, atendiendo a la diversidad lingüística del Ecuador. I.LL.3.1.1.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 15"
        },
        {
            # SEMANA 7 (Semana 2 de la Microcurricular) - DISCURSOS Y JUICIOS DE VALOR (Pág. 19)
            "row_idx": 10,
            "dcd": "LL.3.2.1. Escuchar discursos orales y formular juicios de valor con respecto a su contenido y forma, y participar de manera respetuosa frente a las intervenciones de los demás.",
            "tema": "Tema: Comunicación Oral: Discursos y juicios de valor / Taller: \"Hoy me convertiré en vendedor de sueños\"\n\nMateriales:\nTexto de Lengua 7mo (Páginas 19 a 25)\nPizarra y marcadores de colores\nGuía Docente pág. 13\nCuadernos",
            "estrategias": [
                ("SEMANA 7 (Semana 2 de la Planificación Microcurricular · 12/10 al 16/10/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente invita a observar la escena de la página 19 (\"¿Prepararon ya su discurso para la Semana de la Expresión Oral?\"). Preguntas detonantes: ¿Cómo te sientes cuando hablas en público?, ¿qué es lo más importante al expresar tus ideas ante otras personas?, ¿qué oradores históricos conoces? (ej. Martin Luther King, Manuela Espejo).", False, 10),
                ("Inserción Educación Socioemocional:", True, 10),
                ("Superación del miedo escénico, seguridad personal, respiración diafragmática y empatía hacia los compañeros que exponen sus opiniones.", False, 10),
                ("DUA", True, 10),
                ("Representación: viñeta dialogada del libro pág. 19 y audición del fragmento del discurso \"Yo tengo un sueño\".\nAcción/expresión: lluvia de ideas oral sobre experiencias al hablar ante un auditorio.\nMotivación: identificación con los personajes que expresan dudas y anhelos.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Lectura y análisis del \"Discurso de un niño\" (págs. 20-21). Preguntas de reflexión crítica: ¿Cuál es el mensaje principal del emisor?, ¿qué tono utiliza?, ¿cómo influyen los gestos, la mirada y la modulación de la voz para conmover o convencer al público? Se analiza qué es un juicio de valor respetuoso sin ofender.", False, 10),
                ("DUA", True, 10),
                ("Representación: texto del discurso adaptado y preguntas de comprensión inferencial.\nAcción/expresión: discusión guiada sobre cómo distinguir argumentos de simples opiniones.\nMotivación: valorar la voz infantil en temas de interés social.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("El docente expone la estructura formal de un discurso oral (págs. 22-24):\n1. Introducción (saludo, captación de atención y presentación del tema).\n2. Desarrollo (argumentos fundamentados, anécdotas y evidencias claras).\n3. Conclusión (resumen de la idea central, llamada a la acción y despedida).\nElaboración de un esquema resumen en el cuaderno con las pautas para emitir juicios de valor respetuosos.", False, 10),
                ("DUA", True, 10),
                ("Representación: mapa conceptual en pizarra que desglosa las tres partes del discurso.\nAcción/expresión: registro ordenado en cuadernos y formulación de juicios de valor orales modelo.\nMotivación: preparación para el rol activo de orador escolar.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Desarrollo del Taller de Comunicación Oral \"Hoy me convertiré en vendedor de sueños\" (pág. 25). En parejas, los estudiantes seleccionan un tema social o escolar, redactan un borrador breve con introducción, desarrollo y conclusión, y presentan una disertación de 2 minutos ante la clase aplicando contacto visual y modulación vocal. Coevaluación con escala estimativa.", False, 10),
                ("DUA", True, 10),
                ("Representación: pauta paso a paso de la pág. 25 para organizar la exposición.\nAcción/expresión: presentación oral en parejas y retroalimentación positiva de la clase.\nMotivación: satisfacción de comunicar ideas propias y ser escuchados con respeto.", False, 10),
            ],
            "indicador": "Escucha discursos orales, parafrasea su contenido, formula juicios de valor sobre su forma y participa de manera respetuosa frente a las intervenciones de los demás. I.LL.3.2.1.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 25"
        },
        {
            # SEMANA 8 (Semana 3 de la Microcurricular) - LECTURA CIENTÍFICA Y BIBLIOTECA (Págs. 26 a 37)
            "row_idx": 11,
            "dcd": "LL.3.3.5. Valorar los aspectos de forma y el contenido de un texto científico, a partir de criterios preestablecidos.\n\nLL.3.3.8. Leer con fluidez y entonación en diversos contextos y con diferentes propósitos.",
            "tema": "Tema: Lectura: Leo para acercarme a la ciencia (textos de divulgación científica) y Uso de la biblioteca escolar\n\nMateriales:\nTexto de Lengua 7mo (Páginas 26 a 37)\nPizarra y marcadores\nCuadernos\nFichas de biblioteca escolar",
            "estrategias": [
                ("SEMANA 8 (Semana 3 de la Planificación Microcurricular · 19/10 al 23/10/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente presenta la portada del bloque de lectura (pág. 26) con ilustraciones paleontológicas y fósiles prehistóricos. Preguntas de indagación previa: ¿Qué investiga un científico paleontólogo?, ¿cómo sabemos que existieron seres vivos hace millones de años?, ¿dónde encontramos lecturas de divulgación científica?", False, 10),
                ("Inserción Educación Ambiental y Curiosidad Científica:", True, 10),
                ("Fomento del pensamiento científico, cuidado de los ecosistemas actuales y comprensión de las eras geológicas del planeta Tierra.", False, 10),
                ("DUA", True, 10),
                ("Representación: imágenes de fósiles de la pág. 26 y lectura modelo expresiva del docente.\nAcción/expresión: lluvia de ideas sobre películas y reportajes de ciencia.\nMotivación: asombro y fascinación por la historia natural.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Lectura comentada y secuencial del artículo de divulgación científica de las páginas 28 a 31 (\"Huellas en el tiempo: cazadores de fósiles\"). Análisis de preguntas inferenciales: ¿Qué diferencia a un artículo científico de un cuento de hadas?, ¿qué función cumplen los datos cuantitativos, fechas y fotografías paratextuales?", False, 10),
                ("DUA", True, 10),
                ("Representación: paratextos destacados (títulos, subtítulos, glosario al pie de página).\nAcción/expresión: subrayado de ideas principales e identificación de palabras científicas.\nMotivación: rol activo como lectores críticos de información verídica.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Estructuración del conocimiento en el cuaderno:\n1. Concepto y características del artículo de divulgación científica (claridad, rigor, objetividad).\n2. Estructura textual: Título llamativo, introducción explicativa, cuerpo con evidencias y conclusión.\n3. Normas de uso de la biblioteca escolar y llenado técnico de fichas de préstamo bibliográfico (pág. 37).", False, 10),
                ("DUA", True, 10),
                ("Representación: organizador visual de llaves en la pizarra con tiza/marcadores de color.\nAcción/expresión: transcripción organizada en el cuaderno y diseño de una ficha de biblioteca modelo.\nMotivación: adquisición de hábitos autónomos de investigación escolar.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Resolución de las actividades de comprensión lectora, vocabulario contextual y juicio valorativo de las páginas 34 a 36 del libro escolar. En parejas, eligen un texto de la biblioteca de aula y completan su ficha de registro bibliográfico (pág. 37).", False, 10),
                ("DUA", True, 10),
                ("Representación: cuestionario estructurado del libro y plantilla de ficha bibliográfica.\nAcción/expresión: trabajo colaborativo en parejas y comprobación de respuestas.\nMotivación: enriquecimiento del léxico y satisfacción lectora.", False, 10),
            ],
            "indicador": "Realiza inferencias fundamentales y proyectivo-valorativas, valora los contenidos y aspectos de forma a partir de criterios preestablecidos, al monitorear su comprensión mediante el uso de estrategias cognitivas. I.LL.3.3.2.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 35"
        },
        {
            # SEMANA 9 (Semana 4 de la Microcurricular) - ESCRITURA: NOTA CIENTÍFICA Y EL VERBO (Págs. 38 a 53)
            "row_idx": 12,
            "dcd": "LL.3.4.1. Relatar textos con secuencia lógica, manejo de conectores y coherencia en el uso de la persona y el tiempo verbal.\n\nLL.3.4.6. Autorregular la producción escrita mediante el proceso de planificación, redacción y revisión.\n\nLL.3.4.10. Integrar en las producciones escritas el verbo (accidentes, modos indicativo, subjuntivo e imperativo) y uso del punto y coma (;).",
            "tema": "Tema: Escritura: Escribo una nota científica (planificación y redacción) y Gramática: El verbo (accidentes, modos) y Uso del punto y coma (;)\n\nMateriales:\nTexto de Lengua 7mo (Páginas 38 a 53)\nPizarra y marcadores\nCuadernos de trabajo\nPlantilla de redacción",
            "estrategias": [
                ("SEMANA 9 (Semana 4 de la Planificación Microcurricular · 26/10 al 30/10/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente propone simular la edición de un boletín escolar científico (pág. 38). Observan una muestra de nota científica sobre innovaciones ecológicas. Se plantean interrogantes: ¿Cómo se redacta un texto claro para compartir un descubrimiento?, ¿qué palabras indican las acciones y hallazgos en una investigación?", False, 10),
                ("Inserción Educación para la Seguridad Vial y Movilidad Sostenible:", True, 10),
                ("Propuesta de temática científica: La evolución de los sistemas de transporte ecológico y el respeto al peatón y al ciclista en las ciudades.", False, 10),
                ("DUA", True, 10),
                ("Representación: modelo visual de una nota informativa en pág. 39.\nAcción/expresión: lluvia de ideas de temas de interés.\nMotivación: utilidad práctica del texto para divulgar saberes.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Análisis de la matriz de planificación textual de la pág. 40 (propósito, destinatario, tema y fuentes). Análisis reflexivo de las oraciones: ¿Qué ocurre si cambiamos el tiempo verbal de presente a pasado o futuro? ¿Por qué es necesario usar el punto y coma (;) cuando las oraciones son largas y ya tienen comas internas?", False, 10),
                ("DUA", True, 10),
                ("Representación: oraciones modelo escritas en pizarra comparando tiempos verbales.\nAcción/expresión: deducción de reglas sintácticas mediante participación oral.\nMotivación: comprensión lógica del funcionamiento del idioma.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Construcción guiada en el cuaderno de:\n1. Proceso de escritura: Planificación, redacción del borrador, revisión y publicación (págs. 40-41).\n2. Gramática del Verbo: Raíz y desinencia, accidentes (persona: 1ra, 2da, 3ra; número: singular, plural; tiempo: presente, pretérito, futuro; modos: indicativo, subjuntivo e imperativo) págs. 42-45.\n3. Reglas ortográficas del punto y coma (;) en oraciones yuxtapuestas y enumeraciones complejas (pág. 52).", False, 10),
                ("DUA", True, 10),
                ("Representación: tablas de conjugación cromáticas y mapa conceptual de modos verbales.\nAcción/expresión: elaboración de resúmenes estructurados en cuadernos.\nMotivación: dominar herramientas lingüísticas para redactar con exactitud.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Los estudiantes planifican y redactan el borrador de su nota científica (págs. 46-47) aplicando la conjugación verbal coherente y el uso del punto y coma (;). Realizan coevaluación con la lista de cotejo del texto y resuelven las actividades prácticas de las páginas 50 a 53.", False, 10),
                ("DUA", True, 10),
                ("Representación: plantilla de producción textual estructurada en 3 párrafos.\nAcción/expresión: redacción individual, corrección entre pares y ejercicios gramaticales.\nMotivación: orgullo de publicar un artículo con rigor gramatical.", False, 10),
            ],
            "indicador": "Produce textos informativos utilizando el proceso de escritura, elementos gramaticales apropiados (verbos, modos y tiempos) y signos de puntuación (punto y coma). I.LL.3.6.1. / I.LL.3.6.3.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 45 y página 53"
        },
        {
            # SEMANA 10 (Semana 5 de la Microcurricular) - LITERATURA: DISFRUTO DE LOS MITOS (Págs. 54 a 65)
            "row_idx": 13,
            "dcd": "LL.3.5.1. Reconocer en un texto literario los elementos característicos que le dan sentido (personajes míticos, orígenes del mundo, cosmovisión).",
            "tema": "Tema: Literatura: Disfruto de los mitos (concepto, características, mitos de creación, dioses y héroes)\n\nMateriales:\nTexto de Lengua 7mo (Páginas 54 a 65)\nPizarra y marcadores\nCuadernos\nAntología de mitos andinos y universales",
            "estrategias": [
                ("SEMANA 10 (Semana 5 de la Planificación Microcurricular · 02/11 al 06/11/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente relata el mito de la creación del Sol y la Luna (pág. 54). Preguntas iniciales: ¿Por qué las culturas antiguas crearon relatos sobre seres sobrenaturales?, ¿qué intentaban explicar?, ¿conocen mitos ecuatorianos como el del volcán Chimborazo y la Mama Tungurahua o el mito Cañari de las Guacamayas?", False, 10),
                ("Inserción Educación Cívica e Interculturalidad:", True, 10),
                ("Valoración de la cosmovisión andina, amazónica y universal, promoviendo el respeto por la memoria ancestral de nuestros pueblos originarios.", False, 10),
                ("DUA", True, 10),
                ("Representación: narración oral ambientada e imágenes míticas del texto págs. 54-55.\nAcción/expresión: lluvia de ideas sobre relatos escuchados en la familia.\nMotivación: fascinación por lo fantástico y el misterio del universo.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Lectura compartida y dramatizada de los mitos griegos y andinos seleccionados en las páginas 56 a 61. Análisis grupal: ¿En qué se diferencia un mito de una leyenda histórica o de un cuento común?, ¿qué virtudes o debilidades representan los dioses y héroes?, ¿qué valores humanos transmiten?", False, 10),
                ("DUA", True, 10),
                ("Representación: lectura por roles con entonación expresiva.\nAcción/expresión: cuadro comparativo entre mito, fábula y cuento.\nMotivación: debate de interpretaciones sobre la moral y las fuerzas de la naturaleza.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Sistematización de conceptos en el cuaderno de trabajo:\n1. Definición del Mito como relato sagrado tradicional.\n2. Características (tiempo primordial, dioses, semidioses, fenómenos explicados mediante símbolos).\n3. Clasificación básica: cosmogónicos (origen del mundo), teogónicos (origen de los dioses) y etiológicos (origen de seres y cosas).\nElaboración de un mapa mental ilustrado en los cuadernos.", False, 10),
                ("DUA", True, 10),
                ("Representación: organizador mental en pizarra con ramificaciones visuales.\nAcción/expresión: síntesis teórica en el cuaderno acompañada de un dibujo simbólico.\nMotivación: integración del arte visual con la literatura.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Los estudiantes resuelven los talleres de comprensión, interpretación de personajes y análisis de símbolos de las páginas 62 a 65 del texto de Lengua. Puesta en común de respuestas. (Nota: Lunes 02 y Martes 03 de Noviembre: Feriados por Día de Difuntos e Independencia de Cuenca).", False, 10),
                ("DUA", True, 10),
                ("Representación: preguntas de análisis del libro y guías de discusión.\nAcción/expresión: trabajo individual y socialización en asamblea de aula.\nMotivación: disfrute estético de la lectura literaria.", False, 10),
            ],
            "indicador": "Reconoce en textos de la literatura oral y escrita (mitos) los elementos característicos que les dan sentido; participa en discusiones sobre textos literarios aportando interpretaciones personales. I.LL.3.7.1.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 63"
        },
        {
            # SEMANA 11 (Semana 6 de la Microcurricular) - ESCRITURA CREATIVA Y EVALUACIÓN SUMATIVA (Págs. 66 a 71)
            "row_idx": 14,
            "dcd": "LL.3.5.1. Recrear textos literarios leídos o escuchados mediante el uso de diversos formatos.\n\nLL.3.4.13. Producir textos de acuerdo con la situación comunicativa en formatos variados (historieta mitológica).",
            "tema": "Tema: Escritura creativa: Historieta mitológica y Evaluación sumativa de la Unidad 1\n\nMateriales:\nTexto de Lengua 7mo (Páginas 66 a 71)\nCartulinas, reglas y lápices de colores\nPizarra y marcadores\nCuestionario sumativo de Unidad 1",
            "estrategias": [
                ("SEMANA 11 (Semana 6 de la Planificación Microcurricular · 09/11 al 13/11/2026)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente presenta muestras de historietas y cómics clásicos y contemporáneos (pág. 66). Se exploran saberes previos: ¿Cómo cuenta una historia el lenguaje del cómic?, ¿qué elementos visuales hacen que un cómic sea emocionante de leer?", False, 10),
                ("Inserción Educación Socioemocional y Creatividad:", True, 10),
                ("Desarrollo de la imaginación, trabajo perseverante, expresión de emociones a través del arte y superación de la frustración en la creación gráfica.", False, 10),
                ("DUA", True, 10),
                ("Representación: modelos visuales de tiras cómicas e historietas míticas en págs. 66-67.\nAcción/expresión: análisis oral de viñetas, bocadillos y onomatopeyas.\nMotivación: entusiasmo infantil por el formato de la historieta.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Análisis guiado de los elementos de la historieta (págs. 67-68): viñetas, planos visuales, bocadillos de diálogo (globo de diálogo, susurro, pensamiento, grito), cartelas del narrador y onomatopeyas sonoras. ¿Cómo transformar un mito leído en una secuencia de 4 a 6 viñetas?", False, 10),
                ("DUA", True, 10),
                ("Representación: cuadro tipológico de globos de diálogo y onomatopeyas.\nAcción/expresión: diseño del guion gráfico previo (storyboard) en borrador.\nMotivación: aprendizaje dinámico de la narrativa secuencial gráfica.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Sistematización de las fases para la creación de una historieta mitológica:\n1. Selección del mito y delimitación de personajes principales.\n2. Redacción del guion por viñeta (inicio, nudo, desenlace).\n3. Dibujo de personajes, rotulado de textos y coloreado.\nRepaso de los contenidos clave de la Unidad 1 para la evaluación sumativa (lengua y cultura, discursos, divulgación científica, el verbo y mitos).", False, 10),
                ("DUA", True, 10),
                ("Representación: esquema secuencial del proceso creativo y repaso en pizarra.\nAcción/expresión: bosquejo ordenado en cuaderno y aclaración de dudas.\nMotivación: consolidación de los aprendizajes del bloque.", False, 10),
                ("APLICACIÓN", True, 11),
                ("En parejas o de forma individual, los estudiantes elaboran su Historieta Mitológica final (págs. 68-69) en cartulina y la exponen en el mural del aula (\"Galería Mitológica de 7mo\").\nAplicación de la Evaluación Sumativa estructurada de la Unidad 1 (págs. 70-71) para verificar el dominio de las destrezas priorizadas. (Cierre de la Planificación Microcurricular: 13 de Noviembre de 2026).", False, 10),
                ("DUA", True, 10),
                ("Representación: formato de evaluación sumativa del texto págs. 70-71.\nAcción/expresión: creación artística de la historieta y resolución de la prueba escrita.\nMotivación: celebración del cierre exitoso de la Unidad 1.", False, 10),
            ],
            "indicador": "Recrea textos literarios leídos mediante adaptaciones creativas (historieta) y demuestra el dominio de destrezas comunicativas de la Unidad 1. I.LL.3.7.2. / I.LL.3.6.1.",
            "evaluacion": "Técnica: prueba\nInstrumento: evaluación formativa página 69 y evaluación sumativa páginas 70 y 71"
        },
        {
            # SEMANA 12 (Semana de Proyecto Integrador)
            "row_idx": 15,
            "dcd": "LL.3.4.13. Producir textos de diversa índole de acuerdo con la situación comunicativa, mediante formatos y recursos diversos.\n\nLL.3.2.2. Proponer intervenciones orales estructuradas para sustentar el producto del proyecto interdisciplinar.",
            "tema": "Tema: EJECUCIÓN, CONSOLIDACIÓN Y SOCIALIZACIÓN DEL PROYECTO INTEGRADOR DEL 1ER TRIMESTRE (1 SEMANA)\n\nMateriales:\nGuía oficial del Proyecto Interdisciplinario\nRúbricas ministeriales de evaluación\nMateriales para estands del aula, dípticos y carteles ilustrados",
            "estrategias": [
                ("SEMANA 12 (Semana de Proyecto Integrador · 16/11 al 23/11/2026 · 1 Semana)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("El docente activa la fase culminante del Proyecto Integrador del 1er Trimestre. Los equipos de trabajo revisan las metas planteadas y el borrador de su producto final. Preguntas de orientación: ¿Cómo nuestro producto final responde a la necesidad comunitaria planteada?, ¿qué saberes de Lengua, Matemática, Ciencias y Estudios Sociales articulamos?", False, 10),
                ("Inserción Educación Cívica y Trabajo Colaborativo:", True, 10),
                ("Liderazgo compartido, equidad de género en los roles de equipo, solidaridad y responsabilidad ética en la entrega de compromisos académicos.", False, 10),
                ("DUA", True, 10),
                ("Representación: rúbrica ministerial del proyecto en cartelera.\nAcción/expresión: verificación en equipo mediante lista de chequeo de avances.\nMotivación: entusiasmo por la culminación del proyecto tangible.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Talleres de coevaluación y retroalimentación entre equipos de estudiantes. Análisis reflexivo sobre los desafíos encontrados durante la investigación y cómo se resolvieron. Revisión de la presentación visual y ortografía del producto final.", False, 10),
                ("DUA", True, 10),
                ("Representación: lista de cotejo coevaluativa de 5 criterios.\nAcción/expresión: diálogo constructivo entre pares y sugerencias de mejora.\nMotivación: cultura de retroalimentación formativa y superación continua.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Estructuración del guion formal de exposición del proyecto:\n1. Saludo protocolario e introducción del problema investigado.\n2. Explicación del proceso y articulación de asignaturas.\n3. Demostración práctica del producto final.\n4. Conclusiones y recomendaciones para la comunidad escolar.", False, 10),
                ("DUA", True, 10),
                ("Representación: esquema guía para la sustentación oral.\nAcción/expresión: elaboración de fichas síntesis para los expositores del grupo.\nMotivación: seguridad y dominio del tema frente al público.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Feria de Aula y Socialización del Proyecto Integrador frente a docentes, directivos y compañeros. Cada equipo sustenta su producto aplicando las habilidades de comunicación oral y escrita adquiridas. Entrega oficial del informe y producto final. (Cierre del Proyecto Integrador: 23 de Noviembre de 2026).", False, 10),
                ("DUA", True, 10),
                ("Representación: estands, carteles, trípticos y medios audiovisuales.\nAcción/expresión: exposición oral pública y demostración práctica.\nMotivación: reconocimiento de la comunidad educativa y valoración del esfuerzo colectivo.", False, 10),
            ],
            "indicador": "Publica y socializa los resultados de proyectos integradores mediante textos estructurados e intervenciones orales claras, coherentes y fundamentadas. I.LL.3.6.4.",
            "evaluacion": "• Instrumentos: Rúbrica oficial ministerial del Proyecto Integrador (30% sumativo del proyecto).\n• Actividades: Sustentación oral del producto final y exposición en feria de aula."
        },
        {
            # SEMANA 13 (Semana de Exámenes Trimestrales)
            "row_idx": 16,
            "dcd": "Evaluación sumativa de las destrezas con criterio de desempeño priorizadas trabajadas en la Unidad 1 de Lengua y Literatura para 7mo EGB (LL.3.1.1., LL.3.2.1., LL.3.3.5., LL.3.4.1., LL.3.4.10., LL.3.5.1.).",
            "tema": "Tema: SEMANA DE EVALUACIONES DEL PRIMER TRIMESTRE Y CIERRE ACADÉMICO (1 SEMANA)\n\nMateriales:\nCuestionarios de base estructurada aprobados por Vicerrectorado\nHojas de respuestas y rúbricas de calificación\nRegistros de notas en el Sistema SIEE",
            "estrategias": [
                ("SEMANA 13 (Semana de Exámenes del Primer Trimestre · 24/11 al 30/11/2026 · 1 Semana)", True, 11),
                ("ERCA / DUA", True, 10),
                ("EXPERIENCIA", True, 11),
                ("Generación de un ambiente de serenidad, confianza y concentración en el aula. Dinámica breve de respiración y relajación previa a la entrega de las evaluaciones. El docente explica las instrucciones generales y los tiempos asignados.", False, 10),
                ("Inserción Educación Socioemocional y Ética:", True, 10),
                ("Gestión de la ansiedad frente a evaluaciones, honestidad individual e integridad académica al responder las pruebas sin recurrir al plagio.", False, 10),
                ("DUA", True, 10),
                ("Representación: lectura pausada y modelada de las instrucciones de la prueba.\nAcción/expresión: preguntas aclaratorias de los estudiantes antes de iniciar.\nMotivación: refuerzo de la autoconfianza y mentalidad de crecimiento.", False, 10),
                ("REFLEXIÓN", True, 11),
                ("Los estudiantes realizan una lectura selectiva de los ítems de la evaluación, planificando el orden de resolución y recordando las estrategias cognitivas trabajadas durante el trimestre.", False, 10),
                ("DUA", True, 10),
                ("Representación: formato de examen claro, tipografía legible y espaciado adecuado.\nAcción/expresión: subrayado de palabras clave en las consignas de la prueba.\nMotivación: autonomía en la gestión del tiempo de examen.", False, 10),
                ("CONCEPTUALIZACIÓN", True, 11),
                ("Consolidación y evocación de los aprendizajes desarrollados en la Unidad 1: orígenes y funcionalidad de la escritura, estructura del discurso oral, lectura y valoración de artículos de divulgación científica, conjugación y modos verbales, uso del punto y coma, y características del mito.", False, 10),
                ("DUA", True, 10),
                ("Representación: ítems diversificados (selección múltiple, relación de columnas, completación, análisis sintáctico y producción breve).\nAcción/expresión: resolución estructurada del instrumento sumativo.\nMotivación: oportunidad de evidenciar el progreso académico trimestral.", False, 10),
                ("APLICACIÓN", True, 11),
                ("Desarrollo autónomo de la evaluación sumativa trimestral (del 24 al 30 de noviembre). Entrega de la prueba, revisión inicial de aciertos y errores, retroalimentación individual formativa y consolidación de calificaciones trimestrales (70% formativo + 30% sumativo) para las Juntas de Curso del 01 de Diciembre.", False, 10),
                ("DUA", True, 10),
                ("Representación: escala cuantitativa y cualitativa ministerial visible.\nAcción/expresión: resolución completa del examen y entrega ordenada.\nMotivación: cierre exitoso del 1er Trimestre escolar.", False, 10),
            ],
            "indicador": "Demuestra el dominio de los indicadores de evaluación prioritarios de la Unidad 1 de Lengua y Literatura: I.LL.3.1.1., I.LL.3.2.1., I.LL.3.3.2., I.LL.3.6.1., I.LL.3.6.3., I.LL.3.7.1.",
            "evaluacion": "• Instrumentos: Cuestionario trimestral de base estructurada aprobado por Vicerrectorado (30% sumativo institucional).\n• Actividades: Aplicación de la evaluación trimestral y consolidación del informe de calificaciones."
        }
    ]

    for item in weeks_data:
        r = table.rows[item["row_idx"]]
        
        # Col 0 (DCD)
        r.cells[0].text = item["dcd"]
        
        # Col 2 (Tema y Recursos)
        r.cells[2].text = item["tema"]
        
        # Col 6 (Estrategias ERCA / DUA)
        c6 = r.cells[6]
        c6.text = "" # limpiar
        for text, is_bold, size_pt in item["estrategias"]:
            p = c6.add_paragraph()
            run = p.add_run(text)
            run.bold = is_bold
            run.font.name = "Times New Roman"
            if size_pt:
                run.font.size = Pt(size_pt)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(1)
        
        # Col 13 (Indicador)
        r.cells[13].text = item["indicador"]
        
        # Col 19 (Actividades Evaluativas)
        r.cells[19].text = item["evaluacion"]

    # Guardar documento
    output_filename = "PLANIFICACION_MICROCURRICULAR_7MO_LENGUA_1ER_TRIMESTRE.docx"
    doc.save(output_filename)
    print(f"Documento completo guardado exitosamente: {output_filename}")

if __name__ == "__main__":
    create_7mo_plan()
